#!/usr/bin/env python3
"""Rellena la plantilla oficial de Kahoot con preguntas, colocando las
respuestas correctas en posiciones aleatorias y equilibradas.

Uso:
  python build_kahoot.py --plantilla KahootQuizTemplate.xlsx \
      --json preguntas.json --salida "Mi Kahoot.xlsx" [--semilla 42]
"""
import argparse
import json
import random
import sys
from collections import Counter

from openpyxl import load_workbook

MAX_PREGUNTA = 120
MAX_RESPUESTA = 75
TIEMPOS_VALIDOS = [5, 10, 15, 20, 30, 45, 60, 90, 120, 180, 240]
MAX_PREGUNTAS = 200
MAX_SEGUIDAS = 2


def detectar_columnas(ws):
    """Busca la fila de cabeceras y devuelve (fila, mapa_columnas)."""
    for fila in ws.iter_rows(min_row=1, max_row=30):
        textos = {c.column: str(c.value).strip().lower() for c in fila if c.value}
        col_q = next((k for k, v in textos.items() if v.startswith("question")), None)
        if col_q is None:
            continue
        cols = {"question": col_q}
        for n in range(1, 5):
            col = next((k for k, v in textos.items() if v.startswith(f"answer {n}")), None)
            if col:
                cols[f"a{n}"] = col
        cols["time"] = next((k for k, v in textos.items() if v.startswith("time")), None)
        cols["correct"] = next((k for k, v in textos.items() if v.startswith("correct")), None)
        if all(k in cols for k in ("a1", "a2", "a3", "a4")) and cols["time"] and cols["correct"]:
            return fila[0].row, cols
    # Respaldo: disposición estándar de la plantilla oficial
    print("AVISO: cabeceras no detectadas; se usa la disposición estándar (fila 8, B-H).")
    return 8, {"question": 2, "a1": 3, "a2": 4, "a3": 5, "a4": 6, "time": 7, "correct": 8}


def validar(preguntas):
    errores = []
    if not preguntas:
        errores.append("No hay preguntas")
    if len(preguntas) > MAX_PREGUNTAS:
        errores.append(f"{len(preguntas)} preguntas: la plantilla admite como máximo {MAX_PREGUNTAS}")
    for i, p in enumerate(preguntas, 1):
        correctas = p["correcta"] if isinstance(p["correcta"], list) else [p["correcta"]]
        total = len(correctas) + len(p["incorrectas"])
        if len(p["pregunta"]) > MAX_PREGUNTA:
            errores.append(f"P{i}: enunciado de {len(p['pregunta'])} caracteres (máx {MAX_PREGUNTA})")
        for r in correctas + p["incorrectas"]:
            if len(r) > MAX_RESPUESTA:
                errores.append(f"P{i}: respuesta de {len(r)} caracteres (máx {MAX_RESPUESTA}): {r!r}")
        if not 2 <= total <= 4:
            errores.append(f"P{i}: tiene {total} respuestas (deben ser entre 2 y 4)")
        if p.get("tiempo", 20) not in TIEMPOS_VALIDOS:
            errores.append(f"P{i}: tiempo {p.get('tiempo')} no válido ({TIEMPOS_VALIDOS})")
    return errores


def asignar_posiciones(preguntas, rng):
    """Posición (0-based) de la primera correcta en cada pregunta, equilibrada
    entre columnas y sin más de MAX_SEGUIDAS repeticiones consecutivas."""
    for _ in range(2000):
        uso = Counter()
        posiciones, ok = [], True
        for p in preguntas:
            n = len(p["incorrectas"]) + (len(p["correcta"]) if isinstance(p["correcta"], list) else 1)
            candidatas = list(range(n))
            if len(posiciones) >= MAX_SEGUIDAS and len(set(posiciones[-MAX_SEGUIDAS:])) == 1:
                candidatas = [c for c in candidatas if c != posiciones[-1]] or candidatas
            minimo = min(uso[c] for c in candidatas)
            menos_usadas = [c for c in candidatas if uso[c] <= minimo + 1]
            eleccion = rng.choice(menos_usadas)
            posiciones.append(eleccion)
            uso[eleccion] += 1
        if len(set(posiciones)) > 1 or len(preguntas) == 1:
            return posiciones
    return posiciones


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plantilla", required=True)
    ap.add_argument("--json", required=True)
    ap.add_argument("--salida", required=True)
    ap.add_argument("--semilla", type=int, default=None)
    args = ap.parse_args()

    datos = json.load(open(args.json, encoding="utf-8"))
    preguntas = datos["preguntas"]

    errores = validar(preguntas)
    if errores:
        print("ERRORES — reescribe estos textos y vuelve a ejecutar:")
        print("\n".join(f"  - {e}" for e in errores))
        sys.exit(1)

    rng = random.Random(args.semilla)
    posiciones = asignar_posiciones(preguntas, rng)

    wb = load_workbook(args.plantilla)
    ws = wb.active
    fila_cab, cols = detectar_columnas(ws)
    fila = fila_cab + 1

    # Limpia la fila de ejemplo y cualquier dato previo (conserva formato,
    # numeración de la columna A y validaciones de la plantilla)
    for r in range(fila, max(ws.max_row, fila + len(preguntas)) + 1):
        for c in cols.values():
            ws.cell(row=r, column=c).value = None

    resumen = Counter()
    for i, (p, pos) in enumerate(zip(preguntas, posiciones)):
        correctas = p["correcta"] if isinstance(p["correcta"], list) else [p["correcta"]]
        incorrectas = p["incorrectas"][:]
        rng.shuffle(incorrectas)
        resto = correctas[1:] + incorrectas
        rng.shuffle(resto)
        respuestas = resto[:pos] + [correctas[0]] + resto[pos:]
        idx_correctas = sorted(respuestas.index(c) + 1 for c in correctas)

        r = fila + i
        if cols["question"] > 1 and ws.cell(row=r, column=cols["question"] - 1).value is None:
            # Columna de numeración a la izquierda, si existe en la plantilla
            ws.cell(row=r, column=cols["question"] - 1).value = i + 1
        ws.cell(row=r, column=cols["question"]).value = p["pregunta"]
        for n, texto in enumerate(respuestas, 1):
            ws.cell(row=r, column=cols[f"a{n}"]).value = texto
        ws.cell(row=r, column=cols["time"]).value = p.get("tiempo", 20)
        ws.cell(row=r, column=cols["correct"]).value = (
            idx_correctas[0] if len(idx_correctas) == 1 else ",".join(map(str, idx_correctas))
        )
        for k in idx_correctas:
            resumen[k] += 1

    wb.save(args.salida)

    # Verificación final
    prim = [pos + 1 for pos in posiciones]
    seguidas = max(
        (len(list(g)) for _, g in __import__("itertools").groupby(prim)), default=0
    )
    if len(preguntas) > 3 and len(set(prim)) == 1:
        print("ERROR: todas las correctas en la misma columna.")
        sys.exit(1)
    print(f"OK: {len(preguntas)} preguntas guardadas en {args.salida}")
    print("Correctas por posición: " + " · ".join(f"{k}: {resumen[k]}" for k in sorted(resumen)))
    print(f"Máximo de correctas seguidas en la misma columna: {seguidas}")


if __name__ == "__main__":
    main()
